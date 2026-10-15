"""Write a blueprint and its generated batches as a question paper and an answer book."""

import re
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..models import Batch, Blueprint, Group, Item, Table
from .figures import render

LATIN_FONT = "Times New Roman"
CJK_FONT = "標楷體"
NUMERALS = "一二三四五六七八九十"
CHOICE_KINDS = {"single_choice", "cloze", "reading"}
FIGURE_WIDTH_CM = {"venn": 9, "route": 12, "street_map": 12}
BLANK = "　　"


def _points(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _span(first: int, last: int) -> str:
    return str(first) if first == last else f"{first}–{last}"


def _set_font(run, size: float | None = None, bold: bool = False) -> None:
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = bold or None


def _new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(2)
    _page_number_footer(section)
    return doc


def _page_number_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(paragraph.add_run("第 "), 10)
    run = paragraph.add_run()
    _set_font(run, 10)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._element.append(element)
    _set_font(paragraph.add_run(" 頁"), 10)


def _paragraph(doc, text: str = "", *, size: float | None = None, bold: bool = False,
               indent: float = 0, hanging: float = 0, align=None, keep: bool = False):
    """keep=True holds the paragraph on the same page as the next one."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = keep or None
    if text:
        _set_font(paragraph.add_run(text), size, bold)
    paragraph.paragraph_format.left_indent = Cm(indent + hanging)
    paragraph.paragraph_format.first_line_indent = Cm(-hanging) if hanging else None
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _add_runs(paragraph, text: str, size: float | None = None) -> None:
    """Cloze placeholders {{n}} become underlined blank numbers."""
    for piece in re.split(r"(\{\{\d+\}\})", text):
        if not piece:
            continue
        match = re.fullmatch(r"\{\{(\d+)\}\}", piece)
        run = paragraph.add_run(f"  {match.group(1)}  " if match else piece)
        _set_font(run, size, bold=bool(match))
        run.underline = bool(match) or None


def _header(doc, blueprint: Blueprint, subtitle: str, pending: bool, answer_book: bool) -> None:
    exam, header = blueprint.exam, blueprint.header
    term = NUMERALS[exam.semester - 1]
    _paragraph(doc, f"{header.school_name}　{exam.academic_year} 學年度第{term}學期　{exam.title}",
               size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _paragraph(doc, f"{header.grade} 年級英文科{subtitle}　範圍：{header.scope_text}",
               size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    fields = [label for label, show in (("____年____班", header.show_class_field),
                                         ("座號：______", header.show_seat_field),
                                         ("姓名：____________", header.show_name_field)) if show]
    if fields and not answer_book:
        _paragraph(doc, "　　".join(fields), align=WD_ALIGN_PARAGRAPH.CENTER)
    if pending:
        _paragraph(doc, "※ 本卷含尚未經人工確認（待審）的題目。", size=10,
                   align=WD_ALIGN_PARAGRAPH.CENTER)


def _part_heading(doc, index: int, part, first: int, show_points: bool) -> None:
    """Sections sharing a part (e.g. cloze + reading) sit under one numbered heading."""
    text = f"{NUMERALS[index]}、{part.title}（第 {_span(first, first + part.total_items - 1)} 題"
    if show_points:
        if part.points_per_item is not None:
            text += f"，每題 {_points(part.points_per_item)} 分"
        text += f"，共 {_points(part.total_points)} 分"
    _paragraph(doc, text + "）", bold=True, keep=True).paragraph_format.space_before = Pt(8)


def _choices(doc, item: Item, indent: float, first_line=None) -> None:
    """Lay options out 4, 2 or 1 per line; the first line may continue an existing paragraph."""
    options = [f"({letter}) {getattr(item.choices, letter)}" for letter in "ABCD"]
    longest = max(len(o) for o in options)
    per_line = 4 if longest <= 16 else 2 if longest <= 36 else 1
    width = 17.4 - indent
    for start in range(0, 4, per_line):
        if first_line is not None and start == 0:
            paragraph, text = first_line, "\t"
        else:
            paragraph, text = _paragraph(doc, indent=indent), ""
        stops = paragraph.paragraph_format.tab_stops
        for k in range(per_line):
            stops.add_tab_stop(Cm(indent + width * k / per_line))
        _add_runs(paragraph, text + "\t".join(options[start:start + per_line]))


def _figure(doc, figure) -> None:
    if isinstance(figure, Table):
        _table(doc, figure)
        return
    image = render(figure)
    paragraph = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    paragraph.add_run().add_picture(BytesIO(image), width=Cm(FIGURE_WIDTH_CM[figure.kind]))


def _table(doc, table: Table) -> None:
    _paragraph(doc, table.title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    grid = doc.add_table(rows=len(table.rows) + 1, cols=len(table.columns) + 1)
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [""] + [c.label + (f" ({c.unit})" if c.unit else "") for c in table.columns]
    rows = [labels] + [[row.label, *row.cells] for row in table.rows]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = grid.cell(r, c)
            cell.text = ""
            cell.paragraphs[0].paragraph_format.keep_with_next = True
            _set_font(cell.paragraphs[0].add_run(value), 11, bold=(r == 0 or c == 0))


def _glosses(doc, glossed, indent: float = 0) -> None:
    if glossed:
        notes = "；".join(f"{g.word} {g.meaning_zh}" for g in glossed)
        _paragraph(doc, f"註：{notes}", size=10, indent=indent)


def _question(doc, item: Item, *, in_group: bool, answer_book: bool) -> None:
    pending = "【待審】" if answer_book and item.status == "pending" else ""
    if item.kind in CHOICE_KINDS:
        mark = f"({item.answer.choice})" if answer_book else f"({BLANK})"
        paragraph = _paragraph(doc, hanging=1.9, keep=True)
        _set_font(paragraph.add_run(f"{mark} {item.number}. {pending}"), bold=answer_book)
        if item.kind == "cloze" and in_group:
            _choices(doc, item, indent=1.9, first_line=paragraph)
        else:
            _add_runs(paragraph, item.stem)
            _choices(doc, item, indent=1.9)
    else:
        paragraph = _paragraph(doc, hanging=0.9)
        _set_font(paragraph.add_run(f"{item.number}. {pending}"))
        _add_runs(paragraph, item.stem)
        if item.prompt and item.prompt.instruction:
            _paragraph(doc, f"（{item.prompt.instruction}）", size=11, indent=0.9)
        if answer_book:
            _paragraph(doc, "答：" + " ／ ".join(item.answer.accepted_text), bold=True, indent=0.9)
        elif item.kind == "guided_writing":
            _paragraph(doc, "答：" + "_" * 60, indent=0.9)
    if item.figure is not None:
        _figure(doc, item.figure)
    _glosses(doc, item.glossed, indent=0.9)
    if answer_book:
        _explanation(doc, item)


def _explanation(doc, item: Item) -> None:
    _paragraph(doc, f"解析：{item.explanation.correct_zh}", size=11, indent=0.9)
    for note in item.explanation.distractors:
        _paragraph(doc, f"({note.option}) {note.reason_zh}", size=11, indent=1.4)


def _group(doc, group: Group, *, answer_book: bool) -> None:
    first, last = group.items[0].number, group.items[-1].number
    paragraph = _paragraph(doc, f"【{_span(first, last)} 題】", bold=True, keep=True)
    paragraph.paragraph_format.space_before = Pt(6)
    if answer_book and group.status == "pending":
        _set_font(paragraph.add_run("【待審】"), bold=True)
    _paragraph(doc, group.instruction, size=11, keep=True)
    if group.figure is not None:
        _figure(doc, group.figure)
    for line in (group.passage or "").splitlines():
        if line.strip():
            _add_runs(_paragraph(doc, indent=0.4), line.strip())
    _glosses(doc, group.glossed)
    for item in group.items:
        _question(doc, item, in_group=True, answer_book=answer_book)


def _build(blueprint: Blueprint, batches: list[Batch], *, answer_book: bool) -> Document:
    by_section = {b.section_id: b for b in batches}
    missing = [s.id for s in blueprint.sections if s.id not in by_section]
    if missing:
        raise ValueError(f"missing batches for sections: {', '.join(missing)}")
    for batch in by_section.values():
        batch.validate_against(blueprint)
    pending = any(x.status == "pending" for b in batches for x in (*b.items, *b.groups))

    doc = _new_document()
    _header(doc, blueprint, "答案與解析" if answer_book else "試題", pending, answer_book)
    first = 1
    for index, part in enumerate(blueprint.parts()):
        _part_heading(doc, index, part, first, blueprint.header.show_points)
        for section in part.sections:
            batch = by_section[section.id]
            for group in batch.groups:
                _group(doc, group, answer_book=answer_book)
            for item in batch.items:
                _question(doc, item, in_group=False, answer_book=answer_book)
        first += part.total_items
    return doc


def write_exam(blueprint: Blueprint, batches: list[Batch], paper_path: str | Path,
               answer_path: str | Path) -> None:
    """Write the question paper and the answer book (answers with explanations)."""
    _build(blueprint, batches, answer_book=False).save(paper_path)
    _build(blueprint, batches, answer_book=True).save(answer_path)
