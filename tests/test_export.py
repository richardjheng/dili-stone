"""Word output: a small hand-written exam covering every item kind and drawn figure."""

import pytest
from docx import Document

from dili.export import write_exam
from dili.models import Batch, Blueprint


def _explained(answer: str, correct: str) -> dict:
    return {"correct_zh": correct,
            "distractors": [{"option": o, "reason_zh": f"選項 {o} 不符合題意。"}
                            for o in "ABCD" if o != answer]}


def _choice(number: int, section: str, kind: str, answer: str = "A", **extra) -> dict:
    return {"id": f"q{number}", "section_id": section, "kind": kind, "number": number,
            "stem": extra.pop("stem", f"Pick the best word for question {number}."),
            "choices": {"A": "walking", "B": "walk", "C": "walked", "D": "walks"},
            "answer": {"choice": answer}, "explanation": _explained(answer, f"第 {number} 題考動名詞。"),
            "targets": {"vocab_ids": ["demo-walk"]}, **extra}


def _written(number: int, section: str, kind: str, task: str, answer: str) -> dict:
    return {"id": f"q{number}", "section_id": section, "kind": kind, "number": number,
            "stem": f"Write the answer for question {number}: b_____e.",
            "prompt": {"task": task, "instruction": "依提示作答"},
            "answer": {"accepted_text": [answer]},
            "explanation": {"correct_zh": f"第 {number} 題答案為 {answer}。"},
            "targets": {"vocab_ids": ["demo-walk"]}}


BLUEPRINT = {
    "id": "demo-exam",
    "exam": {"title": "第二次段考", "exam_type": "term", "academic_year": "115", "semester": 1,
             "duration_minutes": 45},
    "header": {"school_name": "示範國中", "grade": 8, "scope_text": "示範範圍 U1–U2"},
    "scope": ["demo-u1"],
    "prerequisite": {"wordlist_id": "moe2000", "lesson_ids": []},
    "sections": [
        {"id": "choice", "title": "單題選擇", "kind": "single_choice", "item_count": 2,
         "points_per_item": "2"},
        {"id": "cloze", "title": "克漏字", "part": "題組", "kind": "cloze", "item_count": 2,
         "points_per_item": "1.5", "group_count": 1, "items_per_group": [2]},
        {"id": "reading", "title": "閱讀測驗", "part": "題組", "kind": "reading", "item_count": 4,
         "points_per_item": "1.5", "group_count": 4, "items_per_group": [1, 1, 1, 1],
         "allow_figures": True},
        {"id": "words", "title": "文意字彙", "kind": "word_in_context", "item_count": 1,
         "points_per_item": "1"},
        {"id": "writing", "title": "依提示作答與翻譯", "kind": "guided_writing", "item_count": 1,
         "points_per_item": "2"},
    ],
    "total_items": 10,
    "total_points": "16",
}

TABLE = {"kind": "table", "id": "fig-menu", "title": "Lunch Menu", "purpose": "prices",
         "columns": [{"id": "price", "label": "Price", "unit": "NT$"}],
         "rows": [{"id": "rice", "label": "Fried rice", "cells": ["80"]},
                  {"id": "soup", "label": "Corn soup", "cells": ["40"]}]}
VENN = {"kind": "venn", "id": "fig-pets", "title": "Two pets", "left_label": "Cats",
        "right_label": "Dogs", "left": ["sleep a lot", "like fish"],
        "right": ["sleep a lot", "like walks"], "intersection": ["sleep a lot"]}
ROUTE = {"kind": "route", "id": "fig-bus", "title": "Bus 7", "service_day": "Saturday",
         "stops": [{"id": "s1", "name": "Park", "arrival_minute": 480, "departure_minute": 480},
                   {"id": "s2", "name": "Zoo", "arrival_minute": 495, "departure_minute": 500}]}
MAP = {"kind": "street_map", "id": "fig-town", "title": "Map of a Town", "width": 2, "height": 1,
       "slots_per_side": 3, "vertical_streets": ["Oak St.", "Pine St.", "Lake St."],
       "horizontal_streets": ["First Ave.", "Second Ave."],
       "places": [{"id": "bank", "name": "Bank", "kind": "building", "block": {"x": 0, "y": 0},
                   "side": "N", "slot": 1},
                  {"id": "bus", "name": "Bus Stop", "kind": "stop", "block": {"x": 0, "y": 0},
                   "side": "N", "slot": 1}],
       "start": {"x": 0, "y": 0}, "facing": "north"}


def _batches() -> list[Batch]:
    reading = [
        {"id": f"g{n}", "section_id": "reading", "kind": "reading", "instruction": "Read and answer.",
         "figure": figure, "items": [_choice(n, "reading", "reading", "B")]}
        for n, figure in ((5, TABLE), (6, VENN), (7, ROUTE), (8, MAP))]
    data = [
        {"run_id": "r1", "section_id": "choice", "produced_by": "test",
         "items": [_choice(1, "choice", "single_choice", glossed=[{"word": "tiny", "meaning_zh": "很小的"}]),
                   _choice(2, "choice", "single_choice", "C")]},
        {"run_id": "r1", "section_id": "cloze", "produced_by": "test",
         "groups": [{"id": "g-cloze", "section_id": "cloze", "kind": "cloze",
                     "instruction": "Choose the best answer for each blank.",
                     "passage": "Tom likes {{3}} in the park.\nHe went {{4}} yesterday.",
                     "items": [_choice(3, "cloze", "cloze"), _choice(4, "cloze", "cloze", "D")]}]},
        {"run_id": "r1", "section_id": "reading", "produced_by": "test", "groups": reading},
        {"run_id": "r1", "section_id": "words", "produced_by": "test",
         "items": [_written(9, "words", "word_in_context", "word_in_context", "believe")]},
        {"run_id": "r1", "section_id": "writing", "produced_by": "test",
         "items": [_written(10, "writing", "guided_writing", "translate_zh_en", "I like to walk.")]},
    ]
    return [Batch.model_validate(d) for d in data]


def _text(path) -> str:
    doc = Document(path)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_writes_paper_and_answer_book(tmp_path):
    paper, answers = tmp_path / "paper.docx", tmp_path / "answers.docx"
    write_exam(Blueprint.model_validate(BLUEPRINT), _batches(), paper, answers)

    paper_text, answer_text = _text(paper), _text(answers)
    assert "示範國中　115 學年度第一學期　第二次段考" in paper_text
    assert "一、單題選擇（第 1–2 題，每題 2 分，共 4 分）" in paper_text
    assert "二、題組（第 3–8 題，每題 1.5 分，共 9 分）" in paper_text
    assert "三、文意字彙（第 9 題，每題 1 分，共 1 分）" in paper_text
    assert "本卷含尚未經人工確認（待審）的題目" in paper_text
    for number in range(1, 11):
        assert f"{number}. " in paper_text
    assert "解析" not in paper_text and "believe" not in paper_text
    assert "(C) 2. 【待審】" in answer_text
    assert "解析：第 9 題答案為 believe。" in answer_text
    assert "(B) 選項 B 不符合題意。" in answer_text
    assert "註：tiny 很小的" in paper_text
    assert "Fried rice" in paper_text

    pictures = Document(paper).inline_shapes
    assert len(pictures) == 3  # Venn, route and map; the table is a native Word table
    assert len(Document(paper).tables) == 1


def test_missing_section_batch_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="missing batches"):
        write_exam(Blueprint.model_validate(BLUEPRINT), _batches()[:-1],
                   tmp_path / "p.docx", tmp_path / "a.docx")
